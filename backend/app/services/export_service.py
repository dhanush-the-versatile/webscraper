"""Export service: render search results to CSV / XLSX / PDF / JSON."""

from __future__ import annotations

import csv
import io
import json
from typing import Any

from sqlalchemy.ext.asyncio import AsyncSession

from app.core.exceptions import NotFoundError, ValidationError
from app.models import ExportFormat, ExportStatus
from app.repositories import (
    ExportRepository,
    SearchRepository,
    SearchResultRepository,
)

_COLUMNS = [
    ("rank", "Rank"),
    ("full_name", "Name"),
    ("headline", "Headline"),
    ("company", "Company"),
    ("country", "Country"),
    ("city", "City"),
    ("seniority", "Seniority"),
    ("years_experience", "Years Exp."),
    ("overall_score", "Score"),
    ("skill_match", "Skill Match"),
    ("technology_match", "Tech Match"),
    ("linkedin_url", "LinkedIn"),
    ("github_url", "GitHub"),
    ("portfolio_url", "Portfolio"),
    ("website_url", "Website"),
    ("summary", "AI Summary"),
]


class ExportService:
    def __init__(self, db: AsyncSession) -> None:
        self.db = db
        self.exports = ExportRepository(db)
        self.searches = SearchRepository(db)
        self.results = SearchResultRepository(db)

    # ------------------------------------------------------------------ #
    async def _rows_for_search(self, search_id: str, user_id: str) -> list[dict[str, Any]]:
        search = await self.searches.get(search_id)
        if search is None or search.user_id != user_id:
            raise NotFoundError("Search not found.")
        results, _ = await self.results.list_for_search(search_id, limit=1000)
        rows: list[dict[str, Any]] = []
        for r in results:
            c = r.candidate
            rows.append(
                {
                    "rank": r.rank,
                    "full_name": c.full_name,
                    "headline": c.headline or "",
                    "company": c.company or "",
                    "country": c.country or "",
                    "city": c.city or "",
                    "seniority": c.seniority.value,
                    "years_experience": c.years_experience if c.years_experience is not None else "",
                    "overall_score": r.overall_score,
                    "skill_match": r.skill_match,
                    "technology_match": r.technology_match,
                    "linkedin_url": c.linkedin_url or "",
                    "github_url": c.github_url or "",
                    "portfolio_url": c.portfolio_url or "",
                    "website_url": c.website_url or "",
                    "summary": r.summary or "",
                }
            )
        return rows

    # ------------------------------------------------------------------ #
    def _to_csv(self, rows: list[dict]) -> bytes:
        buf = io.StringIO()
        writer = csv.writer(buf)
        writer.writerow([label for _, label in _COLUMNS])
        for row in rows:
            writer.writerow([row.get(key, "") for key, _ in _COLUMNS])
        return buf.getvalue().encode("utf-8-sig")  # BOM so Excel opens UTF-8 correctly

    def _to_json(self, rows: list[dict]) -> bytes:
        return json.dumps(rows, indent=2, ensure_ascii=False).encode("utf-8")

    def _to_xlsx(self, rows: list[dict]) -> bytes:
        from openpyxl import Workbook
        from openpyxl.styles import Font

        wb = Workbook()
        ws = wb.active
        ws.title = "Candidates"
        ws.append([label for _, label in _COLUMNS])
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for row in rows:
            ws.append([row.get(key, "") for key, _ in _COLUMNS])
        for idx, (key, _) in enumerate(_COLUMNS, start=1):
            width = max([len(str(r.get(key, ""))) for r in rows[:200]] + [12])
            ws.column_dimensions[ws.cell(row=1, column=idx).column_letter].width = min(width + 2, 50)
        out = io.BytesIO()
        wb.save(out)
        return out.getvalue()

    def _to_pdf(self, rows: list[dict], title: str) -> bytes:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.lib.units import mm
        from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

        buf = io.BytesIO()
        doc = SimpleDocTemplate(
            buf, pagesize=landscape(A4),
            leftMargin=12 * mm, rightMargin=12 * mm, topMargin=14 * mm, bottomMargin=14 * mm,
        )
        styles = getSampleStyleSheet()
        pdf_cols = ["rank", "full_name", "company", "country", "seniority", "overall_score", "skill_match"]
        header = ["#", "Name", "Company", "Country", "Seniority", "Score", "Skill Match"]
        data = [header] + [[str(r.get(c, "")) for c in pdf_cols] for r in rows]
        table = Table(data, repeatRows=1)
        table.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#111827")),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                    ("FONTSIZE", (0, 0), (-1, -1), 8),
                    ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#d1d5db")),
                    ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f9fafb")]),
                    ("LEFTPADDING", (0, 0), (-1, -1), 4),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 4),
                ]
            )
        )
        doc.build([Paragraph(title, styles["Title"]), Spacer(1, 6 * mm), table])
        return buf.getvalue()

    # ------------------------------------------------------------------ #
    async def export_search(self, *, user_id: str, search_id: str, fmt: ExportFormat) -> tuple[bytes, str, str]:
        """Render an export synchronously. Returns (payload, filename, media_type)."""
        rows = await self._rows_for_search(search_id, user_id)
        if not rows:
            raise ValidationError("This search has no results to export.")

        record = await self.exports.create(
            user_id=user_id, search_id=search_id, format=fmt,
            status=ExportStatus.PROCESSING, row_count=len(rows),
        )
        try:
            if fmt == ExportFormat.CSV:
                payload, ext, media = self._to_csv(rows), "csv", "text/csv; charset=utf-8"
            elif fmt == ExportFormat.JSON:
                payload, ext, media = self._to_json(rows), "json", "application/json"
            elif fmt == ExportFormat.XLSX:
                payload, ext, media = (
                    self._to_xlsx(rows), "xlsx",
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            elif fmt == ExportFormat.PDF:
                payload, ext, media = self._to_pdf(rows, "Candidate Export"), "pdf", "application/pdf"
            else:  # pragma: no cover
                raise ValidationError(f"Unsupported export format: {fmt}")
        except Exception as exc:
            await self.exports.update(record, status=ExportStatus.FAILED, error=str(exc))
            raise
        filename = f"candidates-{search_id[:8]}.{ext}"
        await self.exports.update(record, status=ExportStatus.COMPLETED, file_path=filename)
        return payload, filename, media
